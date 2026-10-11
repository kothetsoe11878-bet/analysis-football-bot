from pathlib import Path
from copy import copy
from datetime import datetime
import csv
import re
import unicodedata

from openpyxl import load_workbook


# ============================================================
# SS2 WORKBOOK AUTO UPDATER
# ============================================================

ROOT = Path(".")
SS2_DIR = ROOT / "data" / "ss2"
BRIDGE_FILE = SS2_DIR / "ss2_bridge_current.csv"


# ------------------------------------------------------------
# Workbook groups
# ------------------------------------------------------------

THREE_LEAGUES = {
    "EPL",
    "La_Liga",
    "Bundesliga",
    "Premier League",
    "La Liga",
}

TWO_LEAGUES = {
    "Serie_A",
    "Ligue_1",
    "Serie A",
    "Ligue 1",
}


# ------------------------------------------------------------
# Possible workbook names
# ------------------------------------------------------------

THREE_WORKBOOK_NAMES = [
    "SS2_3Leagues_EPL_LaLiga_Bundesliga_25_26_26_27_SOURCE_SCORE_VERIFIED_FINAL-1.xlsx",
    "SS2_3Leagues_EPL_LaLiga_Bundesliga_25_26_26_27_SOURCE_SCORE_VERIFIED_FINAL.xlsx",
    "Cgpt_3leagues E.B.L_25_26_27_auto final.xlsx",
]

TWO_WORKBOOK_NAMES = [
    "Cgpt_2leagues_SerieA_Ligue1_25_26_27_auto final.xlsx",
    "SS2_SerieA_Ligue1_25_26_26_27_AUTO_FINAL.xlsx",
]


# ============================================================
# BASIC HELPERS
# ============================================================

def clean_text(value):
    if value is None:
        return ""

    return str(value).strip()


def normalize_name(value):
    """
    Normalize team/sheet names for matching.
    Does NOT alter the displayed workbook name.
    """
    value = clean_text(value)

    # Remove favorite arrow
    value = value.replace("↑", "").replace("↓", "")

    # Unicode normalization
    value = unicodedata.normalize("NFKD", value)

    value = "".join(
        c for c in value
        if not unicodedata.combining(c)
    )

    value = value.lower()

    # Common punctuation
    value = value.replace("&", "and")

    # Remove punctuation/spaces
    value = re.sub(r"[^a-z0-9]+", "", value)

    return value


def normalize_league(value):
    value = clean_text(value)

    aliases = {
        "e0": "EPL",
        "england": "EPL",
        "premierleague": "EPL",
        "epl": "EPL",

        "sp1": "La_Liga",
        "laliga": "La_Liga",
        "la liga": "La_Liga",
        "esp1": "La_Liga",

        "d1": "Bundesliga",
        "bundesliga": "Bundesliga",

        "i1": "Serie_A",
        "seriea": "Serie_A",
        "serie a": "Serie_A",

        "f1": "Ligue_1",
        "ligue1": "Ligue_1",
        "ligue 1": "Ligue_1",
    }

    key = value.lower().replace("_", "").strip()

    return aliases.get(key, value)


def normalize_h_a(value):
    value = clean_text(value).upper()

    if value in {"H", "HOME"}:
        return "H"

    if value in {"A", "AWAY"}:
        return "A"

    return value


def strip_arrow(value):
    value = clean_text(value)
    return value.replace("↑", "").replace("↓", "").strip()


# ============================================================
# TEAM ALIASES
# ============================================================

TEAM_ALIASES = {
    # England
    "manchesterunited": "manchesterunited",
    "manunited": "manchesterunited",
    "manutd": "manchesterunited",

    "manchestercity": "manchestercity",
    "mancity": "manchestercity",

    "tottenhamhotspur": "tottenhamhotspur",
    "tottenham": "tottenhamhotspur",
    "spurs": "tottenhamhotspur",

    "westhamunited": "westhamunited",
    "westham": "westhamunited",

    "newcastleunited": "newcastleunited",
    "newcastle": "newcastleunited",

    "wolverhamptonwanderers": "wolverhamptonwanderers",
    "wolves": "wolverhamptonwanderers",

    "brightonandhovealbion": "brightonandhovealbion",
    "brighton": "brightonandhovealbion",

    "nottinghamforest": "nottinghamforest",
    "nottmforest": "nottinghamforest",
    "nottforest": "nottinghamforest",

    "crystalpalace": "crystalpalace",

    "astonvilla": "astonvilla",

    # Spain
    "athleticclub": "athleticclub",
    "athleticbilbao": "athleticclub",

    "atleticomadrid": "atleticomadrid",
    "atleticomad": "atleticomadrid",

    "realbetis": "realbetis",
    "betis": "realbetis",

    "realsociedad": "realsociedad",

    "realmadrid": "realmadrid",

    "barcelona": "barcelona",

    "sevilla": "sevilla",

    "villarreal": "villarreal",

    "celta": "celtavigo",
    "celtavigo": "celtavigo",

    "rayo": "rayovallecano",
    "rayovallecano": "rayovallecano",

    "espanyol": "espanyol",

    "mallorca": "mallorca",

    "girona": "girona",

    "getafe": "getafe",

    "osasuna": "osasuna",

    "alaves": "alaves",

    "valencia": "valencia",

    "levante": "levante",

    "elche": "elche",

    # Italy
    "inter": "inter",
    "internazionale": "inter",
    "intermilan": "inter",

    "milan": "milan",
    "acmilan": "milan",

    "roma": "roma",
    "asroma": "roma",

    "lazio": "lazio",

    "juventus": "juventus",

    "napoli": "napoli",

    "atalanta": "atalanta",

    "fiorentina": "fiorentina",

    "bologna": "bologna",

    # Germany
    "bayernmunich": "bayernmunich",
    "bayern": "bayernmunich",

    "borussiadortmund": "borussiadortmund",
    "dortmund": "borussiadortmund",

    "rbleipzig": "rbleipzig",
    "leipzig": "rbleipzig",

    "eintrachtfrankfurt": "eintrachtfrankfurt",
    "frankfurt": "eintrachtfrankfurt",

    "borussiamonchengladbach": "borussiamonchengladbach",
    "monchengladbach": "borussiamonchengladbach",
    "mgladbach": "borussiamonchengladbach",

    "koln": "koln",
    "cologne": "koln",

    "werderbremen": "werderbremen",

    "stpauli": "stpauli",

    # France
    "parissaintgermain": "parissaintgermain",
    "parisgermain": "parissaintgermain",
    "parissg": "parissaintgermain",
    "psg": "parissaintgermain",

    "monaco": "monaco",

    "lyon": "lyon",
    "olympiquelyon": "lyon",

    "marseille": "marseille",
    "om": "marseille",

    "lille": "lille",

    "nice": "nice",

    "rennes": "rennes",

    "lens": "lens",

    "strasbourg": "strasbourg",
}


def team_key(value):
    key = normalize_name(value)

    return TEAM_ALIASES.get(key, key)


# ============================================================
# DATE
# ============================================================

def format_date(value):
    """
    Convert different date formats to DD/MM/YYYY.
    """

    if value is None:
        return ""

    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")

    text = clean_text(value)

    if not text:
        return ""

    formats = [
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%m/%d/%Y",
        "%Y.%m.%d",
        "%d.%m.%Y",
    ]

    for fmt in formats:
        try:
            dt = datetime.strptime(text, fmt)
            return dt.strftime("%d/%m/%Y")
        except ValueError:
            pass

    # If already looks like DD/MM/YYYY
    if re.match(r"^\d{1,2}/\d{1,2}/\d{4}$", text):
        parts = text.split("/")
        return (
            f"{int(parts[0]):02d}/"
            f"{int(parts[1]):02d}/"
            f"{int(parts[2]):04d}"
        )

    return text


# ============================================================
# HANDICAP DISPLAY
# ============================================================

def format_handicap(value):
    """
    Display:
      0      -> 0
      0.25   -> 0.25
      0.5    -> 0.5
      0.75   -> 0.75
      1      -> 1
      1.25   -> 1.25
    """

    try:
        value = abs(float(value))
    except (TypeError, ValueError):
        return clean_text(value)

    if value == int(value):
        return str(int(value))

    return f"{value:g}"


def get_myanmar_handicap(value):
    """
    Myanmar notation used in the SS2 workbook.
    Based on absolute Asian Handicap line.
    """

    try:
        h = round(abs(float(value)), 2)
    except (TypeError, ValueError):
        return ""

    mapping = {
        0.00: "D",

        0.25: "L-50",
        0.50: "L-100",
        0.75: "1+50",

        1.00: "1D",
        1.25: "1-50",
        1.50: "1-100",
        1.75: "2+50",

        2.00: "2D",
        2.25: "2-50",
        2.50: "2-100",
        2.75: "3+50",

        3.00: "3D",
        3.25: "3-50",
        3.50: "3-100",
        3.75: "4+50",

        4.00: "4D",
    }

    return mapping.get(h, "")


def build_odds_display(ah, myanmar_value=""):
    """
    Workbook display:
       0.5 (L-100)
       1.5 (1-100)
    """

    line = format_handicap(ah)

    myanmar = clean_text(myanmar_value)

    if not myanmar:
        myanmar = get_myanmar_handicap(ah)

    if myanmar:
        return f"{line} ({myanmar})"

    return line


# ============================================================
# BODY RESULT
# ============================================================

def body_cells(result):
    """
    Return:
       Body Win, Body Loss

    Rules:
       WIN + HALF_WIN  -> ✓
       LOSS + HALF_LOSS -> ×
       PUSH -> blank
    """

    result = clean_text(result).upper()

    if result in {"WIN", "HALF_WIN"}:
        return "✓", ""

    if result in {"LOSS", "HALF_LOSS"}:
        return "", "×"

    return "", ""


# ============================================================
# FIND WORKBOOK
# ============================================================

def find_workbook(candidates, required_words):
    """
    First try exact names.
    Then search by required words.
    """

    for name in candidates:
        path = SS2_DIR / name

        if path.exists():
            return path

    xlsx_files = list(SS2_DIR.glob("*.xlsx"))

    for path in xlsx_files:
        lower = path.name.lower()

        if all(word.lower() in lower for word in required_words):
            return path

    return None


# ============================================================
# READ BRIDGE
# ============================================================

def load_bridge():
    if not BRIDGE_FILE.exists():
        raise FileNotFoundError(
            f"Bridge file not found: {BRIDGE_FILE}"
        )

    rows = []

    with BRIDGE_FILE.open(
        "r",
        encoding="utf-8-sig",
        newline=""
    ) as f:

        reader = csv.DictReader(f)

        if not reader.fieldnames:
            raise RuntimeError(
                "ss2_bridge_current.csv has no header."
            )

        for raw in reader:

            row = {
                key.strip(): clean_text(value)
                for key, value in raw.items()
            }

            # Ignore completely empty rows
            if not any(row.values()):
                continue

            rows.append(row)

    return rows


# ============================================================
# HEADER DETECTION
# ============================================================

def find_header_row(ws):
    """
    Find SS2 team-sheet header.

    Expected visible columns:
      L
      Date
      Team
      Odds
      Opponent
      Body (Win)
      Body (Loss)
      H/A
      Score
    """

    max_scan = min(ws.max_row, 15)

    for row_number in range(1, max_scan + 1):

        values = {}

        for col in range(1, min(ws.max_column, 15) + 1):
            value = clean_text(
                ws.cell(row_number, col).value
            )

            if value:
                values[value.lower()] = col

        required = {
            "date",
            "team",
            "odds",
            "opponent",
            "h/a",
            "score",
        }

        if required.issubset(values.keys()):
            return row_number, values

    return None, None


# ============================================================
# COLUMN MAP
# ============================================================

def build_columns(headers):
    """
    Convert header names to actual Excel columns.
    """

    result = {}

    for name, col in headers.items():

        key = name.strip().lower()

        result[key] = col

    aliases = {
        "line": "l",
        "body win": "body (win)",
        "body loss": "body (loss)",
        "ha": "h/a",
    }

    for alias, target in aliases.items():

        if alias in result and target not in result:
            result[target] = result[alias]

    return result


# ============================================================
# EXISTING ROW KEY
# ============================================================

def make_key(date, team, opponent, h_a, score):
    return (
        format_date(date),
        team_key(strip_arrow(team)),
        team_key(strip_arrow(opponent)),
        normalize_h_a(h_a),
        clean_text(score),
    )


# ============================================================
# COPY ROW STYLE
# ============================================================

def copy_row_style(ws, source_row, target_row):
    """
    Copy formatting from previous data row.
    """

    if source_row < 1:
        return

    for col in range(1, ws.max_column + 1):

        source = ws.cell(source_row, col)
        target = ws.cell(target_row, col)

        if source.has_style:
            target._style = copy(source._style)

        if source.number_format:
            target.number_format = source.number_format

        if source.alignment:
            target.alignment = copy(source.alignment)

        if source.protection:
            target.protection = copy(source.protection)


# ============================================================
# TEAM SHEET MAP
# ============================================================

def build_sheet_map(wb):
    """
    Map normalized team names -> worksheet.
    """

    result = {}

    for ws in wb.worksheets:

        header_row, headers = find_header_row(ws)

        if header_row is None:
            continue

        key = team_key(ws.title)

        if key:
            result[key] = ws

    return result


# ============================================================
# WORKBOOK LEAGUE FILTER
# ============================================================

def workbook_accepts_league(workbook_type, league):
    league = normalize_league(league)

    if workbook_type == "three":
        return league in THREE_LEAGUES

    if workbook_type == "two":
        return league in TWO_LEAGUES

    return False


# ============================================================
# UPDATE ONE TEAM SHEET
# ============================================================

def update_team_sheet(ws, bridge_rows):
    """
    Update one team sheet.

    Existing rows are updated.
    Missing rows are appended.
    """

    header_row, headers = find_header_row(ws)

    if header_row is None:
        return 0, 0, 0

    cols = build_columns(headers)

    required_columns = [
        "date",
        "team",
        "odds",
        "opponent",
        "h/a",
        "score",
    ]

    for col_name in required_columns:
        if col_name not in cols:
            print(
                f"  WARNING: {ws.title}: "
                f"missing column '{col_name}'"
            )
            return 0, 0, 0

    # --------------------------------------------------------
    # Existing match index
    # --------------------------------------------------------

    existing = {}

    last_data_row = header_row

    for row_number in range(
        header_row + 1,
        ws.max_row + 1
    ):

        date = ws.cell(
            row_number,
            cols["date"]
        ).value

        team = ws.cell(
            row_number,
            cols["team"]
        ).value

        opponent = ws.cell(
            row_number,
            cols["opponent"]
        ).value

        h_a = ws.cell(
            row_number,
            cols["h/a"]
        ).value

        score = ws.cell(
            row_number,
            cols["score"]
        ).value

        if not any(
            [
                date,
                team,
                opponent,
                h_a,
                score,
            ]
        ):
            continue

        key = make_key(
            date,
            team,
            opponent,
            h_a,
            score,
        )

        existing[key] = row_number

        last_data_row = row_number

    added = 0
    updated = 0
    skipped = 0

    # --------------------------------------------------------
    # Process bridge rows
    # --------------------------------------------------------

    for item in bridge_rows:

        date = format_date(item.get("date"))

        team = clean_text(item.get("team"))
        opponent = clean_text(item.get("opponent"))

        h_a = normalize_h_a(item.get("h_a"))

        arrow = clean_text(item.get("arrow"))

        ah = item.get("ah")

        myanmar = clean_text(
            item.get("myanmar_ah")
        )

        score = clean_text(
            item.get("score")
        )

        body_result = clean_text(
            item.get("body_result")
        )

        if not team or not opponent:
            skipped += 1
            continue

        # ----------------------------------------------------
        # Favorite arrow
        # ----------------------------------------------------

        team_display = team
        opponent_display = opponent

        if arrow:
            team_display = f"↑ {team}"
        else:
            try:
                if float(ah) > 0:
                    opponent_display = f"↑ {opponent}"
            except (TypeError, ValueError):
                pass

        # ----------------------------------------------------
        # Odds
        # ----------------------------------------------------

        odds_display = build_odds_display(
            ah,
            myanmar
        )

        # ----------------------------------------------------
        # Body result
        # ----------------------------------------------------

        body_win, body_loss = body_cells(
            body_result
        )

        # ----------------------------------------------------
        # Match key
        # ----------------------------------------------------

        key = make_key(
            date,
            team_display,
            opponent_display,
            h_a,
            score,
        )

        if key in existing:

            row_number = existing[key]

            # Update values only.
            ws.cell(
                row_number,
                cols["date"]
            ).value = date

            ws.cell(
                row_number,
                cols["team"]
            ).value = team_display

            ws.cell(
                row_number,
                cols["odds"]
            ).value = odds_display

            ws.cell(
                row_number,
                cols["opponent"]
            ).value = opponent_display

            if "body (win)" in cols:
                ws.cell(
                    row_number,
                    cols["body (win)"]
                ).value = body_win

            if "body (loss)" in cols:
                ws.cell(
                    row_number,
                    cols["body (loss)"]
                ).value = body_loss

            ws.cell(
                row_number,
                cols["h/a"]
            ).value = h_a

            ws.cell(
                row_number,
                cols["score"]
            ).value = score

            updated += 1

        else:

            # New row
            last_data_row += 1
            row_number = last_data_row

            # Copy style from previous data row.
            if row_number > header_row + 1:
                copy_row_style(
                    ws,
                    row_number - 1,
                    row_number
                )

            # ------------------------------------------------
            # L / index column
            # ------------------------------------------------

            if "l" in cols:

                numbers = []

                for r in range(
                    header_row + 1,
                    row_number
                ):

                    value = ws.cell(
                        r,
                        cols["l"]
                    ).value

                    try:
                        numbers.append(
                            int(float(value))
                        )
                    except (
                        TypeError,
                        ValueError
                    ):
                        pass

                next_number = (
                    max(numbers) + 1
                    if numbers
                    else 1
                )

                ws.cell(
                    row_number,
                    cols["l"]
                ).value = next_number

            # ------------------------------------------------
            # Write row
            # ------------------------------------------------

            ws.cell(
                row_number,
                cols["date"]
            ).value = date

            ws.cell(
                row_number,
                cols["team"]
            ).value = team_display

            ws.cell(
                row_number,
                cols["odds"]
            ).value = odds_display

            ws.cell(
                row_number,
                cols["opponent"]
            ).value = opponent_display

            if "body (win)" in cols:
                ws.cell(
                    row_number,
                    cols["body (win)"]
                ).value = body_win

            if "body (loss)" in cols:
                ws.cell(
                    row_number,
                    cols["body (loss)"]
                ).value = body_loss

            ws.cell(
                row_number,
                cols["h/a"]
            ).value = h_a

            ws.cell(
                row_number,
                cols["score"]
            ).value = score

            existing[key] = row_number

            added += 1

    return added, updated, skipped


# ============================================================
# UPDATE WORKBOOK
# ============================================================

def update_workbook(path, workbook_type, bridge_rows):

    print()
    print("=" * 70)
    print(f"WORKBOOK: {path.name}")
    print("=" * 70)

    wb = load_workbook(path)

    sheet_map = build_sheet_map(wb)

    # --------------------------------------------------------
    # Filter bridge rows for this workbook
    # --------------------------------------------------------

    accepted_rows = []

    for row in bridge_rows:

        league = normalize_league(
            row.get("league", "")
        )

        if workbook_accepts_league(
            workbook_type,
            league
        ):
            accepted_rows.append(row)

    print(
        f"Bridge rows for this workbook: "
        f"{len(accepted_rows)}"
    )

    # --------------------------------------------------------
    # Group by team
    # --------------------------------------------------------

    grouped = {}

    for row in accepted_rows:

        team = clean_text(
            row.get("team")
        )

        key = team_key(team)

        grouped.setdefault(
            key,
            []
        ).append(row)

    total_added = 0
    total_updated = 0
    total_skipped = 0

    # --------------------------------------------------------
    # Update sheets
    # --------------------------------------------------------

    for team_key_value, rows in grouped.items():

        ws = sheet_map.get(team_key_value)

        if ws is None:

            print(
                f"WARNING: No worksheet found "
                f"for team: {rows[0].get('team')}"
            )

            total_skipped += len(rows)

            continue

        added, updated, skipped = update_team_sheet(
            ws,
            rows
        )

        total_added += added
        total_updated += updated
        total_skipped += skipped

        if added or updated:
            print(
                f"{ws.title}: "
                f"+{added} added, "
                f"{updated} updated"
            )

    # --------------------------------------------------------
    # Save
    # --------------------------------------------------------

    wb.save(path)

    print()
    print(
        f"Saved: {path.name}"
    )

    print(
        f"Added={total_added}, "
        f"Updated={total_updated}, "
        f"Skipped={total_skipped}"
    )

    return (
        total_added,
        total_updated,
        total_skipped,
    )


# ============================================================
# MAIN
# ============================================================

def main():

    print("=" * 70)
    print("SS2 WORKBOOK AUTO UPDATER")
    print("=" * 70)

    if not SS2_DIR.exists():
        raise FileNotFoundError(
            f"Directory not found: {SS2_DIR}"
        )

    bridge_rows = load_bridge()

    print(
        f"Bridge rows loaded: {len(bridge_rows)}"
    )

    # --------------------------------------------------------
    # Find 3-League workbook
    # --------------------------------------------------------

    three_workbook = find_workbook(
        THREE_WORKBOOK_NAMES,
        ["epl", "laliga", "bundesliga"]
    )

    # --------------------------------------------------------
    # Find 2-League workbook
    # --------------------------------------------------------

    two_workbook = find_workbook(
        TWO_WORKBOOK_NAMES,
        ["serie", "ligue"]
    )

    if three_workbook is None:
        print(
            "WARNING: 3-League workbook not found."
        )
    else:
        update_workbook(
            three_workbook,
            "three",
            bridge_rows
        )

    if two_workbook is None:
        print(
            "WARNING: 2-League workbook not found."
        )
    else:
        update_workbook(
            two_workbook,
            "two",
            bridge_rows
        )

    print()
    print("=" * 70)
    print("SS2 WORKBOOK UPDATE COMPLETED")
    print("=" * 70)


if __name__ == "__main__":
    main()
